from flask import Flask, render_template, jsonify, request, redirect, url_for, session
from datetime import datetime
import pandas as pd
import openpyxl
import os
import sys
import time
import traceback
import json

import io
from flask import Flask, render_template, jsonify, request, redirect, url_for, session, send_file

app = Flask(__name__)

app.secret_key = 'mini_project_secret_key'

# --- CONFIGURATION (Start) ---
# Detect environment
IS_CLOUD_RUN = os.environ.get('K_SERVICE') is not None
HAS_SUPABASE = os.environ.get("SUPABASE_URL") is not None

if IS_CLOUD_RUN:
    # Cloud Run (Stateless): Must use /tmp
    FILE_NAME = '/tmp/data.xlsx'
    BASE_DIR = '/tmp'
else:
    # When running as a PyInstaller EXE, __file__ points to the temp extraction
    # folder (sys._MEIPASS) which is wiped on exit. We must store user data files
    # next to the actual .exe, so we use sys.executable's directory instead.
    if getattr(sys, 'frozen', False):
        # Running as packaged EXE
        BASE_DIR = os.path.dirname(os.path.abspath(sys.executable))
    else:
        # Running as normal Python script
        BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    FILE_NAME = os.path.join(BASE_DIR, 'sample_gemini.xlsx')

USE_CLOUD_STORAGE = HAS_SUPABASE
BUCKET_NAME = os.environ.get("BUCKET_NAME", "chitfund-data")

# --- CLOUD STORAGE HELPERS (SUPABASE) ---
def sync_down():
    """Downloads the excel file from Supabase Storage if it exists."""
    if not USE_CLOUD_STORAGE: return
    try:
        from supabase import create_client
        url = os.environ.get("SUPABASE_URL")
        key = os.environ.get("SUPABASE_KEY")
        if not url or not key:
            print("Supabase credentials missing.")
            return

        supabase = create_client(url, key)
        
        # Download file (returns bytes)
        print(f"Downloading data.xlsx from bucket '{BUCKET_NAME}'...")
        response = supabase.storage.from_(BUCKET_NAME).download('data.xlsx')
        
        if response:
            with open(FILE_NAME, 'wb') as f:
                f.write(response)
            print("Download successful.")
        else:
            print("No remote file found.")
            # If Cloud Run and no remote file, we might crash if we don't have a seed.
            # But currently we assume seed is in bucket.
            
    except Exception as e:
        print(f"Supabase Download Error: {e}")

# Initial Sync
if USE_CLOUD_STORAGE: sync_down()

# Cloud Run Fallback: If /tmp/data.xlsx still doesn't exist (sync failed or new setup),
# try to copy from local seed if available in container
if IS_CLOUD_RUN and not os.path.exists(FILE_NAME):
    local_seed = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'sample_gemini.xlsx')
    if os.path.exists(local_seed):
        print("Copying local seed to /tmp...")
        import shutil
        shutil.copy(local_seed, FILE_NAME)

def sync_up():
    """Uploads the local excel file to Supabase Storage."""
    if not USE_CLOUD_STORAGE: return
    try:
        from supabase import create_client
        url = os.environ.get("SUPABASE_URL")
        key = os.environ.get("SUPABASE_KEY")
        if not url or not key: return

        supabase = create_client(url, key)
        
        print(f"Uploading data.xlsx to bucket '{BUCKET_NAME}'...")
        with open(FILE_NAME, 'rb') as f:
            # Upsert is handled via file_options in some versions, or explicit 'upload' with overwriting behavior
            # supabase-py storage upload signature varies. safer to try upload, if fail try update (or use upsert option)
            # Standard py client usually supports file_options={"upsert": "true"}
            supabase.storage.from_(BUCKET_NAME).upload('data.xlsx', f, file_options={"upsert": "true"})
            
        print("Upload successful.")
    except Exception as e:
        # Fallback: if upload fails (maybe file exists and upsert didn't work), try 'update'
        try:
           print(f"Upload failed ({e}), trying update...")
           with open(FILE_NAME, 'rb') as f:
               supabase.storage.from_(BUCKET_NAME).update('data.xlsx', f, file_options={"upsert": "true"})
           print("Update successful.")
        except Exception as e2:
           print(f"Supabase Upload/Update Error: {e2}")

# Initial Sync
if USE_CLOUD_STORAGE: sync_down()

# --- UTILS ---
def clean_num(val):
    try:
        return float(str(val).replace(',', '').replace('₹', '').strip())
    except:
        return 0.0

def clean_plan_amount(plan_str):
    try:
        cleaned = str(plan_str).replace(',', '').replace('₹', '').split('.')[0]
        return int(cleaned)
    except: return 0

def normalize_text(text):
    if not text: return ""
    return str(text).lower().replace('.', '').replace(' ', '').replace(':', '').strip()

# --- PAYMENT RECORDS DB HANDLER ---
PAID_DB_FILE = os.path.join(BASE_DIR, 'payment_records.json')

def load_paid_db():
    if not os.path.exists(PAID_DB_FILE): return {}
    try:
        with open(PAID_DB_FILE, 'r') as f: return json.load(f)
    except: return {}

def save_paid_db(data):
    with open(PAID_DB_FILE, 'w') as f: json.dump(data, f, indent=4)

# --- MEMBER PHONE NUMBERS DB (for WhatsApp receipts) ---
PHONE_DB_FILE = os.path.join(BASE_DIR, 'member_phones.json')

def load_phone_db():
    if not os.path.exists(PHONE_DB_FILE): return {}
    try:
        with open(PHONE_DB_FILE, 'r') as f: return json.load(f)
    except: return {}

def save_phone_db(data):
    with open(PHONE_DB_FILE, 'w') as f: json.dump(data, f, indent=4)

def phone_key(name, area):
    return f"{normalize_text(name)}_{normalize_text(area)}"

# --- HELPER: GET ALL SHEET NAMES (MONTHS) ---
def get_all_sheet_names():
    if not os.path.exists(FILE_NAME): return []
    try:
        wb = openpyxl.load_workbook(FILE_NAME, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except: return []

# --- 1. DASHBOARD READER (Supports specific sheet_name) ---
# --- 1. DASHBOARD READER (Supports specific sheet_name) ---
def get_excel_data(sheet_name=None):
    if not os.path.exists(FILE_NAME): return []
    try:
        # DEBUG: Print sheet names to verify order
        wb_temp = openpyxl.load_workbook(FILE_NAME, read_only=True)
        all_sheets = wb_temp.sheetnames
        print(f"DEBUG: All Sheets: {all_sheets}")
        wb_temp.close()

        # If no sheet specified, default to the FIRST sheet (Index 0)
        # because run_auction_batch moves the NEW sheet to the Front.
        if not sheet_name:
            if all_sheets:
                sheet_name = all_sheets[0] 
            else:
                return []
            
        print(f"DEBUG: Reading sheet '{sheet_name}'")
        df = pd.read_excel(FILE_NAME, sheet_name=sheet_name, header=None, engine='openpyxl')
    except Exception as e:
        print(f"ERROR in get_excel_data: {e}")
        traceback.print_exc()
        return []

    paid_db = load_paid_db()
    members_list = []

    # --- STEP 1: BUILD MASTER LIST ---
    member_map = {}
    name_only_map = {}

    # --- STEP 1: SCAN LEDGER BLOCKS (Authoritative Member Data) ---
    # Ledger columns are typically:
    # A (0): Month/Label, B (1): Plan/Name, C (2): Commission, D (3): Amount
    # F (5): Month/Label, G (6): Plan/Name, H (7): Commission, I (8): Amount
    
    ledger_configs = [
        {'label_col': 0, 'data_col': 1, 'comm_col': 2, 'amt_col': 3},
        {'label_col': 5, 'data_col': 6, 'comm_col': 7, 'amt_col': 8}
    ]

    for config in ledger_configs:
        curr_member = None
        curr_area = "General"
        
        for r in range(df.shape[0]):
            try:
                cell_label = str(df.iloc[r, config['label_col']]).strip().upper()
                
                # Detect Member Header
                if cell_label == 'NAME':
                    raw_name = str(df.iloc[r, config['data_col']]).strip().title()
                    if not raw_name or raw_name.upper() in ['NAN', 'NONE', '']:
                        curr_member = None
                        continue
                        
                    raw_area = str(df.iloc[r, config['comm_col']]).strip().title()
                    if not raw_area or raw_area.upper() in ['NAN', 'NONE', '']:
                        raw_area = "General"
                    
                    norm_n = normalize_text(raw_name)
                    norm_a = normalize_text(raw_area)
                    unique_key = f"{norm_n}_{norm_a}"
                    
                    if unique_key not in member_map:
                        curr_member = {
                            'name': raw_name,
                            'area': raw_area,
                            'total': 0,
                            'items': [],
                            'payment_id': f"{raw_name}_0".replace(" ", ""), # Temporary
                            'is_paid': False,
                            'paid_date': None
                        }
                        members_list.append(curr_member)
                        member_map[unique_key] = curr_member
                        name_only_map.setdefault(norm_n, []).append(curr_member)
                    else:
                        curr_member = member_map[unique_key]

                # Detect Individual Plan Rows (Numeric Month)
                elif curr_member and cell_label.replace('.','',1).isdigit():
                    month_label = str(df.iloc[r, config['label_col']]).strip()
                    plan_val = str(df.iloc[r, config['data_col']]).strip()
                    comm_val = str(df.iloc[r, config['comm_col']]).strip()
                    amt = clean_num(df.iloc[r, config['amt_col']])
                    
                    if amt > 0:
                        # Generate Item ID
                        item_id = f"{normalize_text(curr_member['name'])}_{int(amt)}_{month_label}_{plan_val}".replace(" ", "").replace(".","")
                        
                        # Check payment status
                        db_entry = paid_db.get(item_id)
                        item_is_paid = False
                        item_paid_date = None
                        item_paid_amount = 0
                        
                        if db_entry:
                            if isinstance(db_entry, dict):
                                item_paid_amount = db_entry.get('amount', 0)
                                item_paid_date = db_entry.get('date')
                                item_is_paid = item_paid_amount >= amt
                            else:
                                item_is_paid = True
                                item_paid_amount = amt
                                if isinstance(db_entry, str): item_paid_date = db_entry

                        curr_member['items'].append({
                            'month': month_label,
                            'plan': plan_val,
                            'commission': comm_val,
                            'amount': amt,
                            'id': item_id,
                            'is_paid': item_is_paid,
                            'paid_amount': item_paid_amount,
                            'paid_date': item_paid_date
                        })
                        
                        # Update Member-level payment_id (legacy support)
                        curr_member['payment_id'] = f"{curr_member['name']}_{int(sum(i['amount'] for i in curr_member['items']))}".replace(" ", "")

                # Reset member at block end
                elif 'TOTAL' in cell_label:
                    curr_member = None

            except Exception as e:
                pass

    # --- STEP 2: SCAN FOR COMMISSIONS (Columns B and G specifically for 'Comm.') ---
    # Many users have 'Comm.' rows in their plans. Our Ledger scan above handles them if they have numeric months.
    # If they are listed separately, we'll find them here.
    # Actually, the ledger scan is comprehensive enough. We'll add a specific check for Row-based Grand Totals.

    # --- STEP 4: RE-CALCULATE MEMBER TOTALS with Partial Logic ---
    final_list = []
    
    from datetime import datetime
    for m in members_list:
        if m['items']:
            # Calculate total from items to be safe
            m['total'] = sum(i['amount'] for i in m['items'])
            m['paid_amount'] = sum(i['paid_amount'] for i in m['items'])
            m['is_paid'] = (m['paid_amount'] >= m['total']) and (m['total'] > 0)
            
            # Find the latest payment date
            latest_date_str = None
            latest_date_obj = None
            for i in m['items']:
                date_str = i.get('paid_date')
                if date_str:
                    try:
                        # Format is typically "%d %b, %I:%M %p" e.g. "01 Mar, 08:43 PM"
                        d_obj = datetime.strptime(date_str, "%d %b, %I:%M %p")
                        if not latest_date_obj or d_obj > latest_date_obj:
                            latest_date_obj = d_obj
                            latest_date_str = date_str
                    except Exception:
                        # Fallback for old custom strings
                        latest_date_str = date_str
            m['paid_date'] = latest_date_str

        else:
            # Fallback for members without items
            m['paid_amount'] = m['total'] if m['is_paid'] else 0
            
        final_list.append(m)

    # --- STEP 5: GET EXPLICIT GRAND TOTAL ---
    # The excel has a trusted Grand Total at Row 23, Col 15 (Index 22, 14 0-indexed)
    import math
    grand_total_val = 0
    try:
        # Check specific cell
        raw_gt = df.iloc[22, 14]
        grand_total_val = clean_num(raw_gt)
        if math.isnan(grand_total_val): grand_total_val = 0
        
        print(f"DEBUG: Grand Total from Excel cell [22,14]: {grand_total_val}")
        # Validate order of magnitude (should be > 100k)
        if grand_total_val < 100000:
             # Fallback scan: maybe its somewhere else or sheet is smaller
             print("DEBUG: Grand Total cell too small, scanning column 14 for 'Grand Total' text...")
             for r in range(df.shape[0]):
                 try:
                     label = str(df.iloc[r, 13]).upper()
                     if 'GRAND TOTAL' in label or 'TOTAL' in label:
                         val = clean_num(df.iloc[r, 14])
                         if not math.isnan(val) and val > 100000: 
                             grand_total_val = val
                             break
                 except: pass
    except Exception as e: 
        print(f"DEBUG Error in Grand Total extraction: {e}")
        pass
    
    if math.isnan(grand_total_val): grand_total_val = 0
    # Calculate a sum of all found members for consistency
    calculated_total = sum(m['total'] for m in final_list)
    
    print(f"DEBUG: Returning {len(final_list)} members. Grand Total (Excel): {grand_total_val}, Calculated Total: {calculated_total}")
    
    # Return structure with metadata
    return {
        'members': sorted(final_list, key=lambda x: x['name']),
        'grand_total': grand_total_val,
        'calculated_total': calculated_total
    }

# --- 2. AUCTION READER ---
def get_auction_plans():
    if not os.path.exists(FILE_NAME): return []
    try: df = pd.read_excel(FILE_NAME, header=None, engine='openpyxl')
    except: return []
    plans = []
    for r in range(2, 50):
        try:
            m_val = str(df.iloc[r, 11]).strip()
            p_val_raw = str(df.iloc[r, 12]).strip()
            if not m_val or m_val.upper() == 'NAN': continue
            if 'TOTAL' in str(df.iloc[r, 13]).upper(): break
            p_val = clean_plan_amount(p_val_raw)
            if p_val > 0:
                winning_bid = clean_num(df.iloc[r, 13])
                winner_val = df.iloc[r, 18] if df.shape[1] > 18 else None
                winner_str = str(winner_val).strip() if winner_val is not None and str(winner_val).strip().upper() != 'NAN' else None
                plans.append({
                    'current_month': m_val,
                    'plan_display': p_val_raw,
                    'plan_value': p_val,
                    'winning_bid': winning_bid,
                    'winner_name': winner_str,
                    'id': f"{m_val}_{p_val}"
                })
        except: pass
    return plans

# --- 3. AUCTION UPDATER (FIXED: Month Format + Zero Commission Total) ---
@app.route('/run-auction-batch', methods=['POST'])
def run_auction_batch():
    if 'user' not in session: return redirect(url_for('login_page'))
    try:
        global_date = request.form.get('global_date')
        sheet_name_suffix = request.form.get('sheet_name')

        updates = {}
        # 1. Parse Inputs
        for key, val in request.form.items():
            if key.startswith('bid_for_'):
                unique_id = key.replace('bid_for_', '')
                new_m = request.form.get(f'new_month_for_{unique_id}')
                bid = request.form.get(f'bid_for_{unique_id}')

                # Enforce bid is 0 in the 20th month
                if new_m and (str(new_m).strip() == '20' or str(new_m).strip() == '20.0'):
                    bid = '0'

                # Check if new_m and bid exist (strings "0" are True, empty strings are False)
                if new_m and bid is not None:
                    try:
                        total_bid = float(bid)
                        plan_val = int(unique_id.split('_')[1])
                        dividend = total_bid / 20
                        payable_amount = (plan_val - total_bid) / 20

                        updates[unique_id] = {
                            'new_month': str(new_m).strip(),
                            'new_dividend': int(dividend),
                            'new_payable': int(payable_amount),
                            'plan_val': plan_val,
                            'total_bid': total_bid,
                            'old_month_prefix': unique_id.split('_')[0]
                        }
                    except: continue

        if not updates: return "Error: No Valid Inputs"

        # Helper: Month Matcher
        def is_same_month(m1, m2):
            s1 = str(m1).strip()
            s2 = str(m2).strip()
            if s1 == s2: return True
            try: return abs(float(s1) - float(s2)) < 0.001
            except: return False

        # Load Workbooks
        wb_data = openpyxl.load_workbook(FILE_NAME, data_only=True)
        source_sheet_data = wb_data.active

        wb = openpyxl.load_workbook(FILE_NAME)
        source_sheet = wb.active
        target_sheet = wb.copy_worksheet(source_sheet)
        target_sheet.title = sheet_name_suffix
        wb.move_sheet(target_sheet, offset=-len(wb.sheetnames)+1)
        wb.active = target_sheet

        # Helper: Month Formatter (20 vs 20.0)
        def process_month_value(val):
            try:
                f_val = float(val)
                if f_val.is_integer(): return int(f_val)
                return f_val
            except: return val

        # Map Summary Locations
        summary_locs = {}
        for name_col in [21, 25]:
             for r in range(1, target_sheet.max_row + 1):
                 name = str(target_sheet.cell(row=r, column=name_col).value).strip().title()
                 if name and name.upper() not in ['NAN', 'NAME', 'AMOUNT']:
                     if name not in summary_locs: summary_locs[name] = []
                     summary_locs[name].append({'row': r, 'col': name_col-1})

        # --- UPDATE LOOP 1: Main Ledger ---
        for col_start in [1, 6]:
            curr_member = None
            block_sum = 0      # Sum of Payable Amount
            dividend_sum = 0   # Sum of Commission/Dividend

            for row in range(1, target_sheet.max_row + 1):
                cell_val = str(target_sheet.cell(row=row, column=col_start).value).upper()

                # Header Reset
                if 'NAME' in cell_val:
                    curr_member = str(target_sheet.cell(row=row, column=col_start+1).value).strip().title()
                    target_sheet.cell(row=row, column=col_start+3).value = global_date
                    block_sum = 0
                    dividend_sum = 0 # Reset dividend sum for new block

                # Data Row Processing
                val_check = str(target_sheet.cell(row=row, column=col_start).value).strip()
                if val_check.replace('.','',1).isdigit():
                    curr_m = val_check
                    curr_p = str(target_sheet.cell(row=row, column=col_start+1).value).strip()
                    curr_pv = clean_plan_amount(curr_p)

                    old_pay_val = source_sheet_data.cell(row=row, column=col_start+3).value
                    old_pay = clean_num(old_pay_val)
                    current_item_amt = old_pay

                    # Get existing dividend (for sum calculation if not updated)
                    old_div_val = source_sheet_data.cell(row=row, column=col_start+2).value
                    current_dividend = clean_num(old_div_val)

                    matched = None
                    for uid, upd in updates.items():
                        if upd['plan_val'] == curr_pv and is_same_month(upd['old_month_prefix'], curr_m):
                            matched = upd; break

                    if matched:
                        new_pay = matched['new_payable']
                        diff = new_pay - old_pay

                        fm = process_month_value(matched['new_month'])

                        target_sheet.cell(row=row, column=col_start).value = fm
                        target_sheet.cell(row=row, column=col_start+2).value = matched['new_dividend']
                        target_sheet.cell(row=row, column=col_start+3).value = new_pay

                        current_item_amt = new_pay
                        current_dividend = matched['new_dividend'] # Update dividend for sum

                        # Update Summary Area (Right side of sheet)
                        if curr_member and curr_member in summary_locs:
                            for loc in summary_locs[curr_member]:
                                old_sum_val = source_sheet_data.cell(row=loc['row'], column=loc['col']).value
                                old_sum = clean_num(old_sum_val)
                                target_sheet.cell(row=loc['row'], column=loc['col']).value = int(old_sum) + int(diff)

                    block_sum += current_item_amt
                    dividend_sum += current_dividend # Add to running total

                # Footer Total Update
                if 'TOTAL' in cell_val:
                    target_sheet.cell(row=row, column=col_start+3).value = int(block_sum)
                    # FIX: Explicitly update the Total Commission cell
                    target_sheet.cell(row=row, column=col_start+2).value = int(dividend_sum)
                    block_sum = 0
                    dividend_sum = 0

        # --- UPDATE LOOP 2: Auction List ---
        for r in range(2, 100):
            try:
                p_val_raw = str(source_sheet_data.cell(row=r, column=13).value)
                p_val = clean_plan_amount(p_val_raw)
                m_val = str(source_sheet_data.cell(row=r, column=12).value).strip()

                if p_val > 0:
                    for uid, upd in updates.items():
                        if upd['plan_val'] == p_val and is_same_month(upd['old_month_prefix'], m_val):
                             fm = process_month_value(upd['new_month'])
                             target_sheet.cell(row=r, column=12).value = fm
                             target_sheet.cell(row=r, column=14).value = upd['total_bid']
                             target_sheet.cell(row=r, column=16).value = upd['new_dividend']
                             target_sheet.cell(row=r, column=18).value = upd['new_payable']
                             
                             # Record Winner Name in Column S (19)
                             winner_name = request.form.get(f'winner_for_{uid}')
                             if winner_name:
                                 target_sheet.cell(row=r, column=19).value = winner_name.strip()
            except: pass

        wb.save(FILE_NAME)
        wb_data.close()
        backup_data_to_sql()
        # SYNC TO CLOUD
        if USE_CLOUD_STORAGE: sync_up()
        return redirect(url_for('dashboard'))

    except PermissionError: return "ERROR: Close Excel file!"
    except Exception as e: return f"Error: {e}\n{traceback.format_exc()}"

# --- ROUTES ---
@app.route('/manifest.json')
def serve_manifest():
    return send_file('static/manifest.json', mimetype='application/json')

@app.route('/sw.js')
def serve_sw():
    return send_file('static/sw.js', mimetype='application/javascript')

@app.route('/offline')
def serve_offline():
    return render_template('offline.html')

@app.route('/favicon.ico')
def serve_favicon():
    return send_file('static/icons/icon-192x192.png', mimetype='image/png')

@app.route('/')
def login_page(): return render_template('login.html') if 'user' not in session else redirect(url_for('dashboard'))

@app.route('/login', methods=['POST'])
def login_logic():
    username = request.form.get('username')
    password = request.form.get('password')
    
    # 1. Fallback to hardcoded admin
    if username == 'admin' and password == 'nitesh2025':
        session['user'] = 'admin'
        return redirect(url_for('dashboard'))
        
    # 2. Check SQLite database
    try:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute("SELECT password FROM users WHERE username = ?", (username,))
        row = cursor.fetchone()
        conn.close()
        if row and row[0] == password:
            session['user'] = username
            return redirect(url_for('dashboard'))
    except Exception as e:
        print(f"Login database error: {e}")
        
    return render_template('login.html', error="Invalid Username or Password")

@app.route('/api/signup', methods=['POST'])
def signup_api():
    try:
        data = request.json
        username = data.get('username')
        password = data.get('password')
        if not username or not password:
            return jsonify({'error': 'Missing username or password'}), 400
            
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        try:
            cursor.execute("INSERT INTO users (username, password) VALUES (?, ?)", (username, password))
            conn.commit()
            success = True
            error_msg = None
        except sqlite3.IntegrityError:
            success = False
            error_msg = 'Username already exists'
        conn.close()
        
        if success:
            return jsonify({'success': True})
        return jsonify({'error': error_msg}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/logout')
def logout(): session.pop('user', None); return redirect(url_for('login_page'))

@app.route('/dashboard')
def dashboard():
    if 'user' not in session: return redirect(url_for('login_page'))
    
    # Check if excel file exists and is valid
    has_data = os.path.exists(FILE_NAME)
    if has_data:
        try:
            sheets = get_all_sheet_names()
            if not sheets:
                has_data = False
                default_sheet = "No Data"
            else:
                default_sheet = sheets[0]
        except:
            has_data = False
            default_sheet = "No Data"
    else:
        default_sheet = "No Data"
        
    return render_template('dashboard.html', timestamp=time.time(), active_sheet=default_sheet, has_data=has_data)
@app.route('/members')
def members_page(): return render_template('members.html') if 'user' in session else redirect(url_for('login_page'))
@app.route('/auction')
def auction_page(): return render_template('auction.html', groups=get_auction_plans()) if 'user' in session else redirect(url_for('login_page'))
@app.route('/api/members')
def get_members_api():
    # Optional: allow fetching data for a specific sheet (Month)
    sheet = request.args.get('sheet')
    return jsonify(get_excel_data(sheet)) if 'user' in session else (jsonify({"error": "Unauthorized"}), 401)

@app.route('/api/sheets')
def get_sheets_api():
    return jsonify({'sheets': get_all_sheet_names()}) if 'user' in session else (jsonify({"error": "Unauthorized"}), 401)

# --- EXCEL EDITOR API ---
import re
from openpyxl.utils.cell import coordinate_to_tuple

import re
from openpyxl.utils.cell import coordinate_to_tuple

def evaluate_vlookup(val_str, ws, ws_data):
    pattern = r'=VLOOKUP\((?:CONCATENATE\(([A-Z]+)([0-9]+)\s*,\s*\"\|\"\s*,\s*([A-Z]+)([0-9]+)\)|([A-Z]+)([0-9]+)\s*\&\s*\"\|\"\s*\&\s*([A-Z]+)([0-9]+))\s*,\s*\$?([A-Z]+)\$?([0-9]+):\$?([A-Z]+)\$?([0-9]+)\s*,\s*2\s*,\s*FALSE\)'
    match = re.search(pattern, val_str)
    
    if not match: return val_str
    groups = match.groups()
    
    if groups[0] is not None:
        col1, row1, col2, row2 = groups[0], groups[1], groups[2], groups[3]
    else:
        col1, row1, col2, row2 = groups[4], groups[5], groups[6], groups[7]
        
    start_col, start_row, end_col, end_row = groups[8], groups[9], groups[10], groups[11]
    
    # 1. Get search key from DATA sheet (evaluated values)
    part1 = ws_data[f"{col1}{row1}"].value
    part2 = ws_data[f"{col2}{row2}"].value
    
    # Use fallback to ws if somehow None in ws_data
    if part1 is None: part1 = ws[f"{col1}{row1}"].value
    if part2 is None: part2 = ws[f"{col2}{row2}"].value
    
    if part1 is None: part1 = ""
    if part2 is None: part2 = ""
    
    # Handle float formatting differences (e.g. 1.0 vs 1)
    def format_part(p):
        try:
            f = float(p)
            if f.is_integer(): return str(int(f))
            return str(f)
        except: return str(p).strip()
        
    search_key = f"{format_part(part1)}|{format_part(part2)}"
    
    # 2. Iterate lookup range on DATA sheet
    start_r, start_c = coordinate_to_tuple(f"{start_col}{start_row}")
    end_r, end_c = coordinate_to_tuple(f"{end_col}{end_row}")
    
    for r in range(start_r, end_r + 1):
        # Dynamically evaluate the lookup key from L (12) and M (13) to avoid stale openpyxl cache
        l_val = ws_data.cell(row=r, column=12).value
        m_val = ws_data.cell(row=r, column=13).value
        
        # Fallback to raw ws in case ws_data is None
        if l_val is None: l_val = ws.cell(row=r, column=12).value
        if m_val is None: m_val = ws.cell(row=r, column=13).value
        
        if l_val is not None and m_val is not None:
            lookup_key = f"{format_part(l_val)}|{format_part(m_val)}"
            if lookup_key == search_key:
                res = ws_data.cell(row=r, column=start_c + 1).value
                # Fallback to raw ws in case cached value is not in ws_data
                if res is None: res = ws.cell(row=r, column=start_c + 1).value
                return res if res is not None else ""
            
    return "#N/A" # Default if not found


@app.route('/api/sheet_data')
def get_sheet_data_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    sheet_name = request.args.get('sheet')
    if not sheet_name: return jsonify({'error': 'No sheet name'}), 400
    
    try:
        # Load twice: Once for formulas, once for values
        wb = openpyxl.load_workbook(FILE_NAME)
        wb_data = openpyxl.load_workbook(FILE_NAME, data_only=True)
        
        if sheet_name not in wb.sheetnames: return jsonify({'error': 'Sheet not found'}), 404
        ws = wb[sheet_name]
        ws_data = wb_data[sheet_name]
        
        data = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=False), 1):
            clean_row = []
            for c_idx, cell in enumerate(row, 1):
                cv = cell.value
                if cv is None: clean_row.append("")
                elif isinstance(cv, datetime): clean_row.append(cv.strftime('%Y-%m-%d'))
                else: 
                    val_str = str(cv)
                    if val_str.startswith("=VLOOKUP"):
                        ev = evaluate_vlookup(val_str, ws, ws_data)
                        clean_row.append(ev)
                    else:
                        clean_row.append(val_str)
            data.append(clean_row)
            
        wb.close()
        wb_data.close()
        return jsonify({'data': data})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/delete_sheet', methods=['POST'])
def delete_sheet_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        req = request.json
        sheet_name = req.get('sheet_name')
        if not sheet_name: return jsonify({'error': 'Missing sheet name'}), 400
        
        wb = openpyxl.load_workbook(FILE_NAME)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return jsonify({'error': 'Sheet not found'}), 404
        
        if len(wb.sheetnames) <= 1:
            wb.close()
            return jsonify({'error': 'Cannot delete the last remaining sheet'}), 400
            
        wb.remove(wb[sheet_name])
        wb.save(FILE_NAME)
        wb.close()
        
        # SYNC TO CLOUD
        if USE_CLOUD_STORAGE: sync_up()
        
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/save_sheet_data', methods=['POST'])
def save_sheet_data_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        req = request.json
        sheet_name = req.get('sheet_name')
        changes = req.get('changes')
        
        if not sheet_name or changes is None: return jsonify({'error': 'Missing data or changes'}), 400
        
        wb = openpyxl.load_workbook(FILE_NAME)
        if sheet_name not in wb.sheetnames: return jsonify({'error': 'Sheet not found'}), 404
        ws = wb[sheet_name]
        
        # Process individual cell changes only, preserving other formulas
        for change in changes:
            r_idx = int(change.get('y', 0))
            c_idx = int(change.get('x', 0))
            val = change.get('value', "")
            
            clean_val = val
            if isinstance(val, str):
                val = val.strip()
                if val.replace('.','',1).isdigit():
                    try:
                        if '.' in val: clean_val = float(val)
                        else: clean_val = int(val)
                    except: clean_val = val
                elif val == "":
                    clean_val = None
            
            ws.cell(row=r_idx+1, column=c_idx+1, value=clean_val)
                
        wb.save(FILE_NAME)
        wb.close()
        # SYNC TO CLOUD
        if USE_CLOUD_STORAGE: sync_up()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/toggle-pay', methods=['POST'])
def toggle_pay():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    item_id = request.json.get('id')
    amount = request.json.get('amount') # New parameter
    if not item_id: return jsonify({'error': 'No ID'}), 400
    
    db = load_paid_db()
    now = datetime.now().strftime("%d %b, %I:%M %p")
    
    if amount is not None:
        # Custom amount updated
        amt_float = float(amount)
        if amt_float <= 0:
            if item_id in db: del db[item_id]
        else:
            db[item_id] = {'amount': amt_float, 'date': now}
        save_paid_db(db)
        return jsonify({'success': True, 'paid_on': now})
    else:
        # Legacy toggle logic for backwards compatibility
        current_entry = db.get(item_id)
        if current_entry:
            del db[item_id]
            new_status = False
            paid_on = None
        else:
            db[item_id] = now
            new_status = True
            paid_on = now
        save_paid_db(db)
        return jsonify({'success': True, 'new_status': new_status, 'paid_on': paid_on})

@app.route('/download_excel')
def download_excel():
    if 'user' not in session: return redirect(url_for('login_page'))
    
    # 1. Get processed data
    data_response = get_excel_data()
    if isinstance(data_response, dict):
        members = data_response.get('members', [])
    else:
        members = data_response 
    
    # 2. Create Workbook using openpyxl directly for formatting
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Payment Receipts"
    
    # Styles
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid") # Light Grey
    bold_font = Font(bold=True)
    
    # Column Widths
    ws.column_dimensions['A'].width = 15 # Month / Label
    ws.column_dimensions['B'].width = 25 # Plan / Name
    ws.column_dimensions['C'].width = 25 # Comm / Area
    ws.column_dimensions['D'].width = 20 # Amount
    
    row_idx = 1
    
    for m in members:
        # Sort out date
        display_date = m.get('paid_date', datetime.now().strftime("%d-%b-%Y"))
        if not display_date: display_date = datetime.now().strftime("%d-%b-%Y")
        
        # --- ROW 1: Header Block [Name Label | Name | Area | Date] ---
        # Cell A: Name Label
        c1 = ws.cell(row=row_idx, column=1, value="Name")
        c1.font = bold_font
        c1.border = thin_border
        
        # Cell B: Name Value
        c2 = ws.cell(row=row_idx, column=2, value=m['name'])
        c2.font = bold_font
        c2.border = thin_border
        
        # Cell C: Area Value
        c3 = ws.cell(row=row_idx, column=3, value=m['area'])
        c3.font = bold_font
        c3.border = thin_border
        
        # Cell D: Date
        c4 = ws.cell(row=row_idx, column=4, value=display_date)
        c4.font = bold_font
        c4.alignment = Alignment(horizontal='right')
        c4.border = thin_border
        
        row_idx += 1
        
        # --- ROW 2: Column Headers [Month | Plan | Commission | Amount] ---
        headers = ["Month", "Plan", "Commission", "Amount"]
        for col_num, h_text in enumerate(headers, 1):
            c = ws.cell(row=row_idx, column=col_num, value=h_text)
            c.fill = header_fill
            c.border = thin_border
        row_idx += 1
        
        # --- ROW 3+: Items ---
        items_to_print = m['items'] if m['items'] else [{'month': '-', 'plan': '-', 'commission': '-', 'amount': m['total']}]
        
        for item in items_to_print:
            # Month
            ws.cell(row=row_idx, column=1, value=item.get('month', '-')).border = thin_border
            # Plan
            # Try to format as number if possible
            plan_val = item.get('plan', '-')
            try: plan_val = float(str(plan_val).replace(',','')) 
            except: pass
            c_plan = ws.cell(row=row_idx, column=2, value=plan_val)
            c_plan.border = thin_border
            if isinstance(plan_val, (int, float)): c_plan.number_format = '#,##0.00'
            
            # Commission
            comm_val = item.get('commission', '-')
            try: comm_val = float(str(comm_val).replace(',','')) 
            except: pass
            c_comm = ws.cell(row=row_idx, column=3, value=comm_val)
            c_comm.border = thin_border
            if isinstance(comm_val, (int, float)): c_comm.number_format = '#,##0.00'

            # Amount
            amt_val = item.get('amount', 0)
            c_amt = ws.cell(row=row_idx, column=4, value=amt_val)
            c_amt.border = thin_border
            c_amt.number_format = '#,##0.00'
            
            row_idx += 1
            
        # --- ROW N: Empty spacers (to match image style roughly) ---
        # Image shows some empty grey rows or just space? 
        # Image has 2 empty rows with grey background before Total. Let's add 1 empty row.
        for col_num in range(1, 5):
             ws.cell(row=row_idx, column=col_num).fill = header_fill
             ws.cell(row=row_idx+1, column=col_num).fill = header_fill
        row_idx += 2

        # --- ROW Footer: Total Payable ---
        c_foot_label = ws.cell(row=row_idx, column=1, value="Total Payable")
        c_foot_label.font = bold_font
        c_foot_label.fill = header_fill
        c_foot_label.border = thin_border
        
        # Merge B and C for footer or just leave empty? Image shows total at right.
        ws.cell(row=row_idx, column=2).fill = header_fill
        ws.cell(row=row_idx, column=2).border = thin_border
        ws.cell(row=row_idx, column=3).fill = header_fill
        ws.cell(row=row_idx, column=3).border = thin_border
        
        c_foot_val = ws.cell(row=row_idx, column=4, value=m['total'])
        c_foot_val.font = bold_font
        c_foot_val.fill = header_fill
        c_foot_val.border = thin_border
        c_foot_val.number_format = '#,##0.00'
        
        row_idx += 1
        
        # Spacer between members
        row_idx += 3

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Payment_Receipts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# --- WHATSAPP RECEIPT SENDING (Meta WhatsApp Cloud API) ---
WHATSAPP_TOKEN = os.environ.get("WHATSAPP_TOKEN")
WHATSAPP_PHONE_NUMBER_ID = os.environ.get("WHATSAPP_PHONE_NUMBER_ID")
WHATSAPP_API_VERSION = os.environ.get("WHATSAPP_API_VERSION", "v20.0")

def generate_receipt_image(member, sheet_name=None):
    """Draws a clean receipt card for one member and returns PNG bytes."""
    from PIL import Image, ImageDraw, ImageFont

    width = 700
    row_h = 34
    header_h = 100
    items = member.get('items') or [{'month': '-', 'plan': '-', 'commission': '-', 'amount': member.get('total', 0)}]
    height = header_h + row_h * (len(items) + 1) + 70

    img = Image.new('RGB', (width, height), '#ffffff')
    draw = ImageDraw.Draw(img)

    def font(size, bold=False):
        candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]
        for c in candidates:
            try:
                return ImageFont.truetype(c, size)
            except Exception:
                continue
        return ImageFont.load_default()

    f_title = font(20, True)
    f_head = font(14, True)
    f_body = font(13)
    navy = '#2c3e50'
    grey = '#7f8c8d'
    border = '#e0e6ed'
    header_bg = '#f1f4f8'

    draw.rectangle([0, 0, width - 1, height - 1], outline=border, width=2)

    # Title
    draw.text((20, 15), "Payment Receipt", font=f_title, fill=navy)

    # Header block (Name | Area | Date)
    y = 55
    draw.rectangle([15, y, width - 15, y + 35], fill=header_bg, outline=border)
    draw.text((25, y + 9), f"{member.get('name', '')}", font=f_head, fill=navy)
    draw.text((280, y + 9), f"{member.get('area', '')}", font=f_head, fill=navy)
    date_str = member.get('paid_date') or datetime.now().strftime('%d-%b-%Y')
    draw.text((width - 20 - draw.textlength(date_str, font=f_head), y + 9), date_str, font=f_head, fill=navy)

    # Column headers
    y += 35
    cols = ["Month", "Plan", "Commission", "Amount"]
    col_x = [25, 205, 385, 560]
    draw.rectangle([15, y, width - 15, y + row_h], fill=header_bg, outline=border)
    for cx, c in zip(col_x, cols):
        draw.text((cx, y + 9), c, font=f_head, fill=grey)

    # Item rows
    y += row_h
    for it in items:
        draw.rectangle([15, y, width - 15, y + row_h], outline=border)
        draw.text((col_x[0], y + 8), str(it.get('month', '-')), font=f_body, fill=navy)
        draw.text((col_x[1], y + 8), str(it.get('plan', '-')), font=f_body, fill=navy)
        draw.text((col_x[2], y + 8), str(it.get('commission', '-')), font=f_body, fill=navy)
        amt_str = "{:,.2f}".format(it.get('amount', 0))
        draw.text((width - 25 - draw.textlength(amt_str, font=f_body), y + 8), amt_str, font=f_body, fill=navy)
        y += row_h

    # Total row
    total_str = "{:,.2f}".format(member.get('total', 0))
    draw.rectangle([15, y, width - 15, y + row_h + 5], fill='#f8f9fa', outline=border)
    draw.text((col_x[0], y + 9), "Total Payable", font=f_head, fill=navy)
    draw.text((width - 25 - draw.textlength(total_str, font=f_head), y + 9), total_str, font=f_head, fill=navy)

    out = io.BytesIO()
    img.save(out, format='PNG')
    out.seek(0)
    return out.read()

def whatsapp_send_image(phone, image_bytes, caption=""):
    """Uploads an image to Meta's WhatsApp Cloud API and sends it to `phone`.
    Returns (success: bool, message: str). Requires WHATSAPP_TOKEN and
    WHATSAPP_PHONE_NUMBER_ID environment variables (see WHATSAPP_SETUP.md)."""
    import requests

    if not WHATSAPP_TOKEN or not WHATSAPP_PHONE_NUMBER_ID:
        return False, "WhatsApp is not configured (missing WHATSAPP_TOKEN / WHATSAPP_PHONE_NUMBER_ID env vars)."

    base = f"https://graph.facebook.com/{WHATSAPP_API_VERSION}/{WHATSAPP_PHONE_NUMBER_ID}"
    headers = {"Authorization": f"Bearer {WHATSAPP_TOKEN}"}

    # 1. Upload media
    try:
        files = {'file': ('receipt.png', image_bytes, 'image/png')}
        data = {'messaging_product': 'whatsapp', 'type': 'image/png'}
        r = requests.post(f"{base}/media", headers=headers, files=files, data=data, timeout=30)
        r.raise_for_status()
        media_id = r.json().get('id')
        if not media_id:
            return False, f"Media upload failed: {r.text}"
    except Exception as e:
        return False, f"Media upload error: {e}"

    # 2. Send image message using the uploaded media id
    try:
        clean_phone = str(phone).strip().replace(" ", "").replace("-", "")
        if not clean_phone.startswith("+"):
            clean_phone = "+" + clean_phone.lstrip("+")
        payload = {
            "messaging_product": "whatsapp",
            "recipient_type": "individual",
            "to": clean_phone.replace("+", ""),
            "type": "image",
            "image": {"id": media_id, "caption": caption[:1024]}
        }
        r = requests.post(f"{base}/messages", headers=headers, json=payload, timeout=30)
        if r.status_code >= 400:
            return False, f"Send failed ({r.status_code}): {r.text}"
        return True, "sent"
    except Exception as e:
        return False, f"Send error: {e}"

def find_member(sheet_name, name, area):
    data_response = get_excel_data(sheet_name)
    members = data_response.get('members', []) if isinstance(data_response, dict) else data_response
    for m in members:
        if normalize_text(m.get('name')) == normalize_text(name) and normalize_text(m.get('area')) == normalize_text(area):
            return m
    return None

@app.route('/api/whatsapp/set_phone', methods=['POST'])
def whatsapp_set_phone():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    name, area, phone = data.get('name'), data.get('area'), data.get('phone', '').strip()
    if not name or not area:
        return jsonify({'error': 'name and area required'}), 400
    db = load_phone_db()
    key = phone_key(name, area)
    if phone:
        db[key] = phone
    elif key in db:
        del db[key]
    save_phone_db(db)
    return jsonify({'success': True})

@app.route('/api/whatsapp/send', methods=['POST'])
def whatsapp_send_one():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    name, area, sheet = data.get('name'), data.get('area'), data.get('sheet')
    phone = data.get('phone')
    if not phone:
        phone = load_phone_db().get(phone_key(name, area))
    if not phone:
        return jsonify({'success': False, 'error': 'No phone number saved for this member.'}), 400

    member = find_member(sheet, name, area)
    if not member:
        return jsonify({'success': False, 'error': 'Member not found in this sheet.'}), 404

    img_bytes = generate_receipt_image(member, sheet)
    caption = f"Payment Receipt - {member.get('name')} ({sheet or ''})"
    ok, msg = whatsapp_send_image(phone, img_bytes, caption)
    return jsonify({'success': ok, 'error': None if ok else msg})

@app.route('/api/whatsapp/send_all', methods=['POST'])
def whatsapp_send_all():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    data = request.json or {}
    sheet = data.get('sheet')
    data_response = get_excel_data(sheet)
    members = data_response.get('members', []) if isinstance(data_response, dict) else data_response
    phone_db = load_phone_db()

    results = []
    for m in members:
        phone = phone_db.get(phone_key(m.get('name'), m.get('area')))
        if not phone:
            continue
        img_bytes = generate_receipt_image(m, sheet)
        caption = f"Payment Receipt - {m.get('name')} ({sheet or ''})"
        ok, msg = whatsapp_send_image(phone, img_bytes, caption)
        results.append({'name': m.get('name'), 'phone': phone, 'success': ok, 'error': None if ok else msg})

    return jsonify({'sent': sum(1 for r in results if r['success']), 'total': len(results), 'results': results})

from itertools import groupby

@app.route('/view_receipts')
def view_receipts():
    if 'user' not in session: return redirect(url_for('login_page'))
    
    sheet_names = get_all_sheet_names()
    # Accept explicit sheet param, else default to latest (index 0)
    target_sheet = request.args.get('sheet')
    if not target_sheet and sheet_names:
        target_sheet = sheet_names[0]
    
    data_response = get_excel_data(target_sheet)
    # Handle dict response
    if isinstance(data_response, dict):
        members = data_response.get('members', [])
    else:
        members = data_response
    
    # Sort by Area (primary) and Name (secondary)
    def sort_key(x):
        return (str(x.get('area', '')).strip().lower(), str(x.get('name', '')).strip().lower())
    
    members.sort(key=sort_key)

    phone_db = load_phone_db()
    for m in members:
        m['phone'] = phone_db.get(phone_key(m.get('name'), m.get('area')), '')
    
    # Continuous Layout: Return flat list, sorted by Area
    return render_template('receipt_preview.html', 
                           members=members, 
                           now=datetime.now(),
                           sheet_names=sheet_names,
                           active_sheet=target_sheet)

# --- REPORT ROUTES ---
@app.route('/reports')
def reports_page():
    if 'user' not in session: return redirect(url_for('login_page'))
    sheets = get_all_sheet_names()
    return render_template('reports.html', sheets=sheets)

@app.route('/excel_editor')
def excel_editor_page():
    return render_template('excel_editor.html') if 'user' in session else redirect(url_for('login_page'))

# --- MEMBER MANAGEMENT EXCEL HELPERS ---
def edit_member_in_excel(old_name, old_area, new_name, new_area):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # 1. Update main ledger blocks (Columns A-D / 1-4 and F-I / 6-9)
            for col_start in [1, 6]:
                for r in range(1, ws.max_row + 1):
                    cell_val = str(ws.cell(row=r, column=col_start).value).strip().upper()
                    if cell_val == 'NAME':
                        curr_name = str(ws.cell(row=r, column=col_start+1).value).strip()
                        curr_area = str(ws.cell(row=r, column=col_start+2).value).strip()
                        if normalize_text(curr_name) == normalize_text(old_name) and normalize_text(curr_area) == normalize_text(old_area):
                            ws.cell(row=r, column=col_start+1, value=new_name.strip())
                            ws.cell(row=r, column=col_start+2, value=new_area.strip())

            # 2. Update summary tables (Col U/21 and Y/25 for Name, Col V/22 and Z/26 for Area)
            for name_col, area_col in [(21, 22), (25, 26)]:
                for r in range(1, ws.max_row + 1):
                    curr_name = str(ws.cell(row=r, column=name_col).value).strip()
                    curr_area = str(ws.cell(row=r, column=area_col).value).strip()
                    if normalize_text(curr_name) == normalize_text(old_name) and normalize_text(curr_area) == normalize_text(old_area):
                        ws.cell(row=r, column=name_col, value=new_name.strip())
                        ws.cell(row=r, column=area_col, value=new_area.strip())
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in edit_member_in_excel: {e}")
        return False

def delete_member_from_excel(name, area):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # 1. Clear ledger block
            for col_start in [1, 6]:
                r = 1
                while r <= ws.max_row:
                    cell_val = str(ws.cell(row=r, column=col_start).value).strip().upper()
                    if cell_val == 'NAME':
                        curr_name = str(ws.cell(row=r, column=col_start+1).value).strip()
                        curr_area = str(ws.cell(row=r, column=col_start+2).value).strip()
                        if normalize_text(curr_name) == normalize_text(name) and normalize_text(curr_area) == normalize_text(area):
                            # Clear cells in Name row
                            ws.cell(row=r, column=col_start).value = None
                            ws.cell(row=r, column=col_start+1).value = None
                            ws.cell(row=r, column=col_start+2).value = None
                            ws.cell(row=r, column=col_start+3).value = None
                            
                            row_runner = r + 1
                            while row_runner <= ws.max_row:
                                runner_label = str(ws.cell(row=row_runner, column=col_start).value).strip().upper()
                                for c in range(col_start, col_start + 4):
                                    ws.cell(row=row_runner, column=c).value = None
                                if 'TOTAL' in runner_label:
                                    break
                                row_runner += 1
                            r = row_runner
                    r += 1

            # 2. Clear from summary tables
            for name_col, area_col, amt_col in [(21, 22, 20), (25, 26, 24)]:
                for r in range(1, ws.max_row + 1):
                    curr_name = str(ws.cell(row=r, column=name_col).value).strip()
                    curr_area = str(ws.cell(row=r, column=area_col).value).strip()
                    if normalize_text(curr_name) == normalize_text(name) and normalize_text(curr_area) == normalize_text(area):
                        ws.cell(row=r, column=name_col).value = None
                        ws.cell(row=r, column=area_col).value = None
                        ws.cell(row=r, column=amt_col).value = None
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in delete_member_from_excel: {e}")
        return False

def add_member_to_excel(name, area):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Find the bottom of columns A-D (Ledger 1)
            last_row_a = 0
            for r in range(ws.max_row, 0, -1):
                if ws.cell(row=r, column=1).value is not None:
                    last_row_a = r
                    break
            if last_row_a == 0:
                last_row_a = ws.max_row
                
            start_row = last_row_a + 3
            
            # Write new ledger block
            ws.cell(row=start_row, column=1, value="Name")
            ws.cell(row=start_row, column=2, value=name.strip())
            ws.cell(row=start_row, column=3, value=area.strip())
            ws.cell(row=start_row, column=4, value=datetime.now().strftime("%Y-%m-%d"))
            
            ws.cell(row=start_row+1, column=1, value="Month")
            ws.cell(row=start_row+1, column=2, value="Plan")
            ws.cell(row=start_row+1, column=3, value="Commission")
            ws.cell(row=start_row+1, column=4, value="Amount")
            
            ws.cell(row=start_row+2, column=1, value="Total Payable")
            ws.cell(row=start_row+2, column=2, value=None)
            ws.cell(row=start_row+2, column=3, value=0)
            ws.cell(row=start_row+2, column=4, value=0)
            
            # Find empty slot in summary columns U/Y
            summary_placed = False
            for r in range(3, ws.max_row + 10):
                val_u = ws.cell(row=r, column=21).value
                val_t = ws.cell(row=r, column=20).value
                if val_u is None and val_t is None:
                    ws.cell(row=r, column=21, value=name.strip())
                    ws.cell(row=r, column=22, value=area.strip())
                    ws.cell(row=r, column=20, value=f"=D{start_row+2}")
                    summary_placed = True
                    break
            
            if not summary_placed:
                for r in range(3, ws.max_row + 10):
                    val_y = ws.cell(row=r, column=25).value
                    val_x = ws.cell(row=r, column=24).value
                    if val_y is None and val_x is None:
                        ws.cell(row=r, column=25, value=name.strip())
                        ws.cell(row=r, column=26, value=area.strip())
                        ws.cell(row=r, column=24, value=f"=D{start_row+2}")
                        summary_placed = True
                        break
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in add_member_to_excel: {e}")
        return False

# --- MEMBER MANAGEMENT API ROUTES ---
@app.route('/api/add_member', methods=['POST'])
def add_member_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        name = data.get('name')
        area = data.get('area')
        if not name or not area:
            return jsonify({'error': 'Missing name or area'}), 400
            
        success = add_member_to_excel(name, area)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to add member to Excel'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/edit_member', methods=['POST'])
def edit_member_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        old_name = data.get('old_name')
        old_area = data.get('old_area')
        new_name = data.get('new_name')
        new_area = data.get('new_area')
        if not old_name or not old_area or not new_name or not new_area:
            return jsonify({'error': 'Missing details'}), 400
            
        success = edit_member_in_excel(old_name, old_area, new_name, new_area)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to edit member in Excel'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/delete_member', methods=['POST'])
def delete_member_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        name = data.get('name')
        area = data.get('area')
        if not name or not area:
            return jsonify({'error': 'Missing details'}), 400
            
        success = delete_member_from_excel(name, area)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to delete member from Excel'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/upload_excel', methods=['POST'])
def upload_excel_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file part'}), 400
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'No selected file'}), 400
        if file and file.filename.endswith('.xlsx'):
            temp_path = os.path.join(BASE_DIR, "temp_upload.xlsx")
            file.save(temp_path)
            
            # Temporarily swap FILE_NAME to parse it
            global FILE_NAME
            original_filename = FILE_NAME
            
            try:
                FILE_NAME = temp_path
                data_res = get_excel_data()
                plans = get_auction_plans()
            finally:
                FILE_NAME = original_filename
                
            # Filter unique member names for selection dropdown
            members = sorted(list(set([m['name'] for m in data_res.get('members', []) if m.get('name')])))
            
            return jsonify({
                'success': True,
                'needs_initial_setup': True,
                'plans': plans,
                'members': members,
                'temp_filename': temp_path
            })
        else:
            return jsonify({'error': 'Invalid format. Please upload an .xlsx file.'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/upload_excel_by_path', methods=['POST'])
def upload_excel_by_path_api():
    """Desktop EXE mode: accepts a local file path from pywebview's native file dialog."""
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        file_path = data.get('file_path')
        if not file_path:
            return jsonify({'error': 'No file_path provided'}), 400
        if not os.path.exists(file_path):
            return jsonify({'error': 'File not found: ' + file_path}), 400
        if not file_path.lower().endswith('.xlsx'):
            return jsonify({'error': 'Invalid format. Please use an .xlsx file.'}), 400

        import shutil
        temp_path = os.path.join(BASE_DIR, 'temp_upload.xlsx')
        shutil.copy2(file_path, temp_path)

        global FILE_NAME
        original_filename = FILE_NAME
        try:
            FILE_NAME = temp_path
            data_res = get_excel_data()
            plans = get_auction_plans()
        finally:
            FILE_NAME = original_filename

        members = sorted(list(set([m['name'] for m in data_res.get('members', []) if m.get('name')])))

        return jsonify({
            'success': True,
            'needs_initial_setup': True,
            'plans': plans,
            'members': members,
            'temp_filename': temp_path
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/submit_initial_setup', methods=['POST'])
def submit_initial_setup_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        temp_filename = data.get('temp_filename', 'temp_upload.xlsx')
        setup_list = data.get('setup', [])
        
        if not os.path.exists(temp_filename):
            return jsonify({'error': 'Uploaded file not found.'}), 400
            
        # Update columns N (winning bid) and S (winner name) in the temp file
        wb = openpyxl.load_workbook(temp_filename)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for s_item in setup_list:
                try:
                    m_target = float(s_item.get('month'))
                    p_target = int(s_item.get('plan_amount'))
                    winner = s_item.get('winner_name')
                    w_bid = float(s_item.get('winning_bid', 0.0))
                    
                    # Scan L (12) and M (13) in plans table
                    for r in range(3, 100):
                        val_m = ws.cell(row=r, column=12).value
                        val_p = ws.cell(row=r, column=13).value
                        if val_m is not None and val_p is not None:
                            try:
                                if abs(float(val_m) - m_target) < 0.001 and int(clean_plan_amount(val_p)) == p_target:
                                    ws.cell(row=r, column=14, value=w_bid)
                                    if winner and winner != 'Unspecified' and winner != '':
                                        ws.cell(row=r, column=19, value=winner)
                                    else:
                                        ws.cell(row=r, column=19, value=None)
                            except: pass
                except: pass
                
        # Save as main filename
        wb.save(FILE_NAME)
        wb.close()
        
        # Clean up temp file
        try: os.remove(temp_filename)
        except: pass
        
        # Sync SQL & Cloud
        backup_data_to_sql()
        if USE_CLOUD_STORAGE: sync_up()
        
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Helper: Add a new plan to the auction plan table
def add_plan_to_excel(month_val, plan_amt_val):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # 1. Find the TOTAL row in Column N (14)
            total_row = None
            for r in range(1, 100):
                val = ws.cell(row=r, column=14).value
                if val == 'TOTAL':
                    total_row = r
                    break
            
            if not total_row:
                total_row = 21
                
            # 2. Shift columns L (12) to R (18) down by 1 row
            for r in range(total_row + 15, total_row - 1, -1):
                for c in range(12, 19):
                    ws.cell(row=r+1, column=c, value=ws.cell(row=r, column=c).value)
                    
            # 3. Write the new plan details at total_row
            ws.cell(row=total_row, column=12, value=float(month_val))
            ws.cell(row=total_row, column=13, value=int(plan_amt_val))
            ws.cell(row=total_row, column=14, value=0)
            ws.cell(row=total_row, column=15, value=f"=L{total_row}&\"|\"&M{total_row}")
            ws.cell(row=total_row, column=16, value=f"=N{total_row}/20")
            ws.cell(row=total_row, column=17, value=f"=L{total_row}&\"|\"&M{total_row}")
            ws.cell(row=total_row, column=18, value=f"=(M{total_row}-N{total_row})/20")
            
            # 4. Update the grand total formula at total_row + 3
            ws.cell(row=total_row+3, column=15, value=f"=SUM(O{total_row+1},R{total_row+1})")
            
            # 5. Dynamically update VLOOKUP ranges in all formulas across the sheet
            import re
            new_end_row = total_row
            for r in range(1, ws.max_row + 1):
                for c in [3, 4, 8, 9]:
                    val = ws.cell(row=r, column=c).value
                    if isinstance(val, str) and 'VLOOKUP' in val:
                        val = re.sub(r'\$O\$3:\$P\$[0-9]+', f'$O$3:$P${new_end_row}', val)
                        val = re.sub(r'\$Q\$3:\$R\$[0-9]+', f'$Q$3:$R${new_end_row}', val)
                        ws.cell(row=r, column=c, value=val)
                        
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in add_plan_to_excel: {e}")
        return False

@app.route('/api/add_plan', methods=['POST'])
def add_plan_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        month = data.get('month')
        plan_amount = data.get('plan_amount')
        if not month or not plan_amount:
            return jsonify({'error': 'Missing month or plan amount'}), 400
            
        success = add_plan_to_excel(month, plan_amount)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to add plan to Excel'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Helper: Edit an existing plan in the spreadsheet
def edit_plan_in_excel(old_month, old_plan, new_month, new_plan):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for r in range(3, 100):
                val_m = ws.cell(row=r, column=12).value
                val_p = ws.cell(row=r, column=13).value
                if val_m is not None and val_p is not None:
                    try:
                        if abs(float(val_m) - float(old_month)) < 0.001 and int(clean_plan_amount(val_p)) == int(old_plan):
                            ws.cell(row=r, column=12, value=float(new_month))
                            ws.cell(row=r, column=13, value=int(new_plan))
                            ws.cell(row=r, column=15, value=f"=L{r}&\"|\"&M{r}")
                            ws.cell(row=r, column=17, value=f"=L{r}&\"|\"&M{r}")
                            ws.cell(row=r, column=18, value=f"=(M{r}-N{r})/20")
                    except: pass
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in edit_plan_in_excel: {e}")
        return False

# Helper: Delete a plan row from the plan table
def delete_plan_from_excel(month, plan):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for r in range(3, 100):
                val_m = ws.cell(row=r, column=12).value
                val_p = ws.cell(row=r, column=13).value
                if val_m is not None and val_p is not None:
                    try:
                        if abs(float(val_m) - float(month)) < 0.001 and int(clean_plan_amount(val_p)) == int(plan):
                            for c in range(12, 20): # Clear L to S
                                ws.cell(row=r, column=c, value=None)
                    except: pass
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in delete_plan_from_excel: {e}")
        return False

# Helper: Add a due/plan line to a member block
def add_member_due_to_excel(member_name, member_area, month, plan):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Find member ledger block
            col_start = None
            r_name = None
            for cs in [1, 6]:
                for r in range(1, ws.max_row + 1):
                    cell_val = str(ws.cell(row=r, column=cs).value).strip().upper()
                    if cell_val == 'NAME':
                        curr_name = str(ws.cell(row=r, column=cs+1).value).strip()
                        curr_area = str(ws.cell(row=r, column=cs+2).value).strip()
                        if normalize_text(curr_name) == normalize_text(member_name) and normalize_text(curr_area) == normalize_text(member_area):
                            col_start = cs
                            r_name = r
                            break
                if col_start: break
                
            if not col_start: continue
            
            # Find the TOTAL row of this member block
            r_total = None
            for r in range(r_name + 1, ws.max_row + 20):
                cell_val = str(ws.cell(row=r, column=col_start).value).strip().upper()
                if 'TOTAL' in cell_val:
                    r_total = r
                    break
                    
            if not r_total: continue
            
            # Shift columns manually down by 1 row
            col_range = range(1, 5) if col_start == 1 else range(6, 10)
            for r_idx in range(ws.max_row + 5, r_total - 1, -1):
                for c_idx in col_range:
                    ws.cell(row=r_idx+1, column=c_idx, value=ws.cell(row=r_idx, column=c_idx).value)
                    
            # Write new due row
            col_letter_m = 'A' if col_start == 1 else 'F'
            col_letter_p = 'B' if col_start == 1 else 'G'
            
            ws.cell(row=r_total, column=col_start, value=float(month))
            ws.cell(row=r_total, column=col_start+1, value=int(plan))
            ws.cell(row=r_total, column=col_start+2, value=f"=VLOOKUP({col_letter_m}{r_total}&\"|\"&{col_letter_p}{r_total},$O$3:$P$100,2,FALSE)")
            ws.cell(row=r_total, column=col_start+3, value=f"=VLOOKUP({col_letter_m}{r_total}&\"|\"&{col_letter_p}{r_total},$Q$3:$R$100,2,FALSE)")
            
            # Update shifted total row
            col_letter_a = 'D' if col_start == 1 else 'I'
            ws.cell(row=r_total+1, column=col_start+3, value=f"=SUM({col_letter_a}{r_name+2}:{col_letter_a}{r_total})")
            
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in add_member_due_to_excel: {e}")
        return False

# Helper: Delete a due/plan line from a member block
def delete_member_due_from_excel(member_name, member_area, month, plan):
    if not os.path.exists(FILE_NAME): return False
    try:
        wb = openpyxl.load_workbook(FILE_NAME)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Find member ledger block
            col_start = None
            r_name = None
            for cs in [1, 6]:
                for r in range(1, ws.max_row + 1):
                    cell_val = str(ws.cell(row=r, column=cs).value).strip().upper()
                    if cell_val == 'NAME':
                        curr_name = str(ws.cell(row=r, column=cs+1).value).strip()
                        curr_area = str(ws.cell(row=r, column=cs+2).value).strip()
                        if normalize_text(curr_name) == normalize_text(member_name) and normalize_text(curr_area) == normalize_text(member_area):
                            col_start = cs
                            r_name = r
                            break
                if col_start: break
                
            if not col_start: continue
            
            # Find the due row inside this block
            r_due = None
            r_total = None
            for r in range(r_name + 1, ws.max_row + 20):
                cell_label = str(ws.cell(row=r, column=col_start).value).strip()
                if 'TOTAL' in cell_label.upper():
                    r_total = r
                    break
                if cell_label.replace('.','',1).isdigit():
                    try:
                        p_val = clean_plan_amount(ws.cell(row=r, column=col_start+1).value)
                        if abs(float(cell_label) - float(month)) < 0.001 and int(p_val) == int(plan):
                            r_due = r
                    except: pass
                    
            if not r_due or not r_total: continue
            
            # Shift columns manually up by 1 row
            col_range = range(1, 5) if col_start == 1 else range(6, 10)
            for r_idx in range(r_due, ws.max_row + 5):
                for c_idx in col_range:
                    ws.cell(row=r_idx, column=c_idx, value=ws.cell(row=r_idx+1, column=c_idx).value)
                    
            # Update shifted total row
            col_letter_a = 'D' if col_start == 1 else 'I'
            if r_total - 2 >= r_name + 2:
                ws.cell(row=r_total-1, column=col_start+3, value=f"=SUM({col_letter_a}{r_name+2}:{col_letter_a}{r_total-2})")
            else:
                ws.cell(row=r_total-1, column=col_start+3, value=0)
                
        wb.save(FILE_NAME)
        wb.close()
        return True
    except Exception as e:
        print(f"Error in delete_member_due_from_excel: {e}")
        return False

# --- ADDITIONAL MEMBER & PLAN API ROUTES ---
@app.route('/api/edit_plan', methods=['POST'])
def edit_plan_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        old_month = data.get('old_month')
        old_plan = data.get('old_plan')
        new_month = data.get('new_month')
        new_plan = data.get('new_plan')
        if not old_month or not old_plan or not new_month or not new_plan:
            return jsonify({'error': 'Missing details'}), 400
            
        success = edit_plan_in_excel(old_month, old_plan, new_month, new_plan)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to edit plan'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/delete_plan', methods=['POST'])
def delete_plan_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        month = data.get('month')
        plan = data.get('plan')
        if not month or not plan:
            return jsonify({'error': 'Missing details'}), 400
            
        success = delete_plan_from_excel(month, plan)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to delete plan'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/add_member_due', methods=['POST'])
def add_member_due_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        member_name = data.get('name')
        member_area = data.get('area')
        month = data.get('month')
        plan = data.get('plan')
        if not member_name or not member_area or not month or not plan:
            return jsonify({'error': 'Missing details'}), 400
            
        success = add_member_due_to_excel(member_name, member_area, month, plan)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to add due line'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/delete_member_due', methods=['POST'])
def delete_member_due_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data = request.json
        member_name = data.get('name')
        member_area = data.get('area')
        month = data.get('month')
        plan = data.get('plan')
        if not member_name or not member_area or not month or not plan:
            return jsonify({'error': 'Missing details'}), 400
            
        success = delete_member_due_from_excel(member_name, member_area, month, plan)
        if success:
            backup_data_to_sql()
            if USE_CLOUD_STORAGE: sync_up()
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Failed to delete due line'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# --- SQL DATA BACKUP SYSTEM ---
import sqlite3

DB_FILE = os.path.join(BASE_DIR, 'backup.db')

def init_db():
    try:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS users (
                username TEXT PRIMARY KEY,
                password TEXT NOT NULL
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS members_backup (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sheet_name TEXT,
                member_name TEXT,
                area TEXT,
                month REAL,
                plan_amount INTEGER,
                commission REAL,
                payable_amount INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS plans_backup (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sheet_name TEXT,
                month REAL,
                plan_amount INTEGER,
                winning_bid REAL,
                commission REAL,
                payable_amount INTEGER,
                winner_name TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Error in init_db: {e}")

init_db()

def backup_data_to_sql():
    if not os.path.exists(FILE_NAME): return False
    try:
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        
        # Create Tables
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS members_backup (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sheet_name TEXT,
                member_name TEXT,
                area TEXT,
                month REAL,
                plan_amount INTEGER,
                commission REAL,
                payable_amount INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS plans_backup (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sheet_name TEXT,
                month REAL,
                plan_amount INTEGER,
                winning_bid REAL,
                commission REAL,
                payable_amount INTEGER,
                winner_name TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # Clear existing backup data to write fresh sync
        cursor.execute('DELETE FROM members_backup')
        cursor.execute('DELETE FROM plans_backup')
        
        # Load workbook
        wb = openpyxl.load_workbook(FILE_NAME, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # --- 1. Parse Members ---
            members_data = []
            for col_start in [1, 6]:
                curr_member = None
                curr_area = None
                for r in range(1, ws.max_row + 1):
                    val_col0 = ws.cell(row=r, column=col_start).value
                    val_str = str(val_col0).strip().upper() if val_col0 is not None else ""
                    
                    if 'NAME' in val_str:
                        curr_member = str(ws.cell(row=r, column=col_start+1).value).strip()
                        curr_area = str(ws.cell(row=r, column=col_start+2).value).strip()
                        continue
                        
                    if 'TOTAL' in val_str:
                        curr_member = None
                        curr_area = None
                        continue
                        
                    if curr_member and val_col0 is not None:
                        # Parse month
                        cell_label = str(val_col0).strip()
                        if cell_label.replace('.','',1).isdigit():
                            try:
                                month_val = float(cell_label)
                                plan_val = int(clean_plan_amount(ws.cell(row=r, column=col_start+1).value))
                                comm_val = float(ws.cell(row=r, column=col_start+2).value or 0.0)
                                pay_val = int(ws.cell(row=r, column=col_start+3).value or 0)
                                members_data.append((sheet_name, curr_member, curr_area, month_val, plan_val, comm_val, pay_val))
                            except: pass
                            
            # Insert members
            if members_data:
                cursor.executemany('''
                    INSERT INTO members_backup (sheet_name, member_name, area, month, plan_amount, commission, payable_amount)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                ''', members_data)
                
            # --- 2. Parse Plans ---
            plans_data = []
            for r in range(3, 100):
                val_m = ws.cell(row=r, column=12).value
                val_p = ws.cell(row=r, column=13).value
                if val_m is not None and val_p is not None:
                    try:
                        month_val = float(val_m)
                        plan_val = int(clean_plan_amount(val_p))
                        bid_val = float(ws.cell(row=r, column=14).value or 0.0)
                        comm_val = float(ws.cell(row=r, column=16).value or 0.0)
                        pay_val = int(ws.cell(row=r, column=18).value or 0)
                        winner_val = ws.cell(row=r, column=19).value # Column S
                        winner_str = str(winner_val).strip() if winner_val else None
                        plans_data.append((sheet_name, month_val, plan_val, bid_val, comm_val, pay_val, winner_str))
                    except: pass
                    
            # Insert plans
            if plans_data:
                cursor.executemany('''
                    INSERT INTO plans_backup (sheet_name, month, plan_amount, winning_bid, commission, payable_amount, winner_name)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                ''', plans_data)
                
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print(f"Error in backup_data_to_sql: {e}")
        return False

@app.route('/api/trigger_backup', methods=['POST'])
def trigger_backup_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        success = backup_data_to_sql()
        if success:
            return jsonify({'success': True})
        else:
            return jsonify({'error': 'Backup execution failed'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ai_suggestions')
def ai_suggestions_api():
    if 'user' not in session: return jsonify({'error': 'Unauthorized'}), 401
    try:
        data_res = get_excel_data()
        members = data_res.get('members', [])
        plans = get_auction_plans()
        
        suggestions = []
        for p in plans:
            m_val = p['current_month']
            p_val = p['plan_value']
            winner = p['winner_name']
            winning_bid = p['winning_bid']
            target_payout = p_val - winning_bid
            
            # Find all member due lines assigned to this plan
            total_members = 0
            paid_count = 0
            collected_amount = 0
            pending_amount = 0
            
            for m in members:
                for item in m.get('items', []):
                    try:
                        item_month = str(item.get('month')).strip()
                        item_plan = int(item.get('plan', 0))
                        
                        # Match current month and plan value
                        if item_plan == p_val and item_month == str(m_val).strip():
                            total_members += 1
                            item_amt = item.get('amount', 0)
                            if item.get('is_paid'):
                                paid_count += 1
                                collected_amount += item_amt
                            else:
                                pending_amount += item_amt
                    except: pass
                    
            # Compute AI suggestion text
            if not winner:
                recommendation = "AI Suggestion: No bid winner has been selected yet. Run the auction first to calculate winning bids and record the winner."
                collected_pct = 0
            else:
                collected_pct = round((collected_amount / (collected_amount + pending_amount) * 100)) if (collected_amount + pending_amount) > 0 else 0
                if pending_amount == 0:
                    recommendation = f"AI Recommendation: 100% of member installments (₹{collected_amount:,}) have been collected. It is fully safe to pay the prize money of ₹{target_payout:,} to the winner {winner}."
                elif collected_amount == 0:
                    recommendation = f"AI Recommendation: No collection has been received yet for this plan. We suggest holding the payout of ₹{target_payout:,} for {winner} until members begin making payments."
                else:
                    recommendation = f"AI Recommendation: You have collected ₹{collected_amount:,} ({collected_pct}%) from paid members. We suggest making a partial payout of ₹{collected_amount:,} to {winner} now, and releasing the remaining ₹{pending_amount:,} as pending members clear their dues."
                    
            suggestions.append({
                'month': m_val,
                'plan_amount': p_val,
                'winning_bid': winning_bid,
                'target_payout': target_payout,
                'winner_name': winner or "Unspecified",
                'collected_amount': collected_amount,
                'pending_amount': pending_amount,
                'paid_count': paid_count,
                'total_count': total_members,
                'collected_pct': collected_pct,
                'recommendation': recommendation
            })
            
        return jsonify({'success': True, 'suggestions': suggestions})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
